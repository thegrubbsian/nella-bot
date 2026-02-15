"""Tests for document text extraction (PDF, DOCX, XLSX)."""

from pathlib import Path

import pytest

from src.tools.extractors import MAX_EXTRACTED_CHARS, extract_text


# ---------------------------------------------------------------------------
# Helpers — create minimal valid files using the same libraries
# ---------------------------------------------------------------------------


def _make_pdf(path: Path, pages: list[str]) -> Path:
    """Create a minimal PDF with the given page texts."""
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(str(path))
    doc.close()
    return path


def _make_docx(path: Path, paragraphs: list[str]) -> Path:
    """Create a minimal DOCX with the given paragraphs."""
    import docx

    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _make_xlsx(path: Path, sheets: dict[str, list[list]]) -> Path:
    """Create a minimal XLSX with named sheets and row data."""
    import openpyxl

    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    wb.close()
    return path


# ---------------------------------------------------------------------------
# PDF extraction
# ---------------------------------------------------------------------------


async def test_extract_pdf_single_page(tmp_path: Path) -> None:
    pdf = _make_pdf(tmp_path / "test.pdf", ["Hello from PDF"])
    result = await extract_text(pdf)
    assert result is not None
    assert "Hello from PDF" in result


async def test_extract_pdf_multi_page(tmp_path: Path) -> None:
    pdf = _make_pdf(tmp_path / "multi.pdf", ["Page one", "Page two"])
    result = await extract_text(pdf)
    assert result is not None
    assert "Page one" in result
    assert "Page two" in result


async def test_extract_pdf_blank_page(tmp_path: Path) -> None:
    """A PDF with only blank pages should return None."""
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()  # blank page, no text inserted
    path = tmp_path / "blank.pdf"
    doc.save(str(path))
    doc.close()

    result = await extract_text(path)
    assert result is None


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------


async def test_extract_docx(tmp_path: Path) -> None:
    docx_path = _make_docx(tmp_path / "test.docx", ["Hello from DOCX", "Second paragraph"])
    result = await extract_text(docx_path)
    assert result is not None
    assert "Hello from DOCX" in result
    assert "Second paragraph" in result


async def test_extract_docx_empty(tmp_path: Path) -> None:
    """A DOCX with only empty paragraphs should return None."""
    docx_path = _make_docx(tmp_path / "empty.docx", ["", "   "])
    result = await extract_text(docx_path)
    assert result is None


# ---------------------------------------------------------------------------
# XLSX extraction
# ---------------------------------------------------------------------------


async def test_extract_xlsx_single_sheet(tmp_path: Path) -> None:
    xlsx = _make_xlsx(tmp_path / "test.xlsx", {"Data": [["Name", "Age"], ["Alice", 30]]})
    result = await extract_text(xlsx)
    assert result is not None
    assert "Alice" in result
    assert "30" in result
    # Single sheet should NOT have a sheet header
    assert "## Sheet:" not in result


async def test_extract_xlsx_multi_sheet(tmp_path: Path) -> None:
    xlsx = _make_xlsx(
        tmp_path / "multi.xlsx",
        {
            "Sales": [["Q1", 100], ["Q2", 200]],
            "Expenses": [["Rent", 500]],
        },
    )
    result = await extract_text(xlsx)
    assert result is not None
    assert "## Sheet: Sales" in result
    assert "## Sheet: Expenses" in result
    assert "100" in result
    assert "500" in result


async def test_extract_xlsx_empty_rows_skipped(tmp_path: Path) -> None:
    xlsx = _make_xlsx(tmp_path / "sparse.xlsx", {"Sheet1": [["data"], [None, None], ["more"]]})
    result = await extract_text(xlsx)
    assert result is not None
    # The empty row (None, None) should be skipped
    lines = [line for line in result.split("\n") if line.strip()]
    assert len(lines) == 2


# ---------------------------------------------------------------------------
# Unsupported / corrupted / edge cases
# ---------------------------------------------------------------------------


async def test_unsupported_format_returns_none(tmp_path: Path) -> None:
    jpg = tmp_path / "photo.jpg"
    jpg.write_bytes(b"\xff\xd8\xff\xe0")
    result = await extract_text(jpg)
    assert result is None


async def test_corrupted_pdf_returns_none(tmp_path: Path) -> None:
    bad_pdf = tmp_path / "bad.pdf"
    bad_pdf.write_bytes(b"not a real pdf")
    result = await extract_text(bad_pdf)
    assert result is None


async def test_corrupted_docx_returns_none(tmp_path: Path) -> None:
    bad_docx = tmp_path / "bad.docx"
    bad_docx.write_bytes(b"not a real docx")
    result = await extract_text(bad_docx)
    assert result is None


async def test_corrupted_xlsx_returns_none(tmp_path: Path) -> None:
    bad_xlsx = tmp_path / "bad.xlsx"
    bad_xlsx.write_bytes(b"not a real xlsx")
    result = await extract_text(bad_xlsx)
    assert result is None


async def test_unknown_extension_returns_none(tmp_path: Path) -> None:
    unknown = tmp_path / "data.zzz"
    unknown.write_bytes(b"whatever")
    result = await extract_text(unknown)
    assert result is None


async def test_truncation_at_max_chars(tmp_path: Path) -> None:
    """Text longer than MAX_EXTRACTED_CHARS gets truncated."""
    # Use DOCX for reliable large-text generation (PDF insert_text clips to page width)
    long_para = "A" * 5000
    num_paragraphs = (MAX_EXTRACTED_CHARS // 5000) + 5
    docx_path = _make_docx(tmp_path / "long.docx", [long_para] * num_paragraphs)
    result = await extract_text(docx_path)
    assert result is not None
    assert result.endswith("[truncated — extracted text too large]")
    # The truncated result should be slightly longer than MAX_EXTRACTED_CHARS
    # due to the appended truncation message
    assert len(result) < MAX_EXTRACTED_CHARS + 100
